from __future__ import annotations

import json
import os
import socket
import time
import traceback
import uuid
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook

from .budget import BudgetUnavailable
from .catalog import cleanup_expired
from .config import settings
from .db import connect, init_db, now, transaction
from .extraction import extract_document
from .provider import embed_texts, embed_visual_inputs, provider_settings
from .retrieval import delete_document_vectors, upsert_text_vectors, upsert_visual_vectors


WORKER_ID = f"{socket.gethostname()}:{os.getpid()}:{uuid.uuid4().hex[:8]}"


def claim_job():
    stamp = now()
    with connect() as db, transaction(db, immediate=True):
        db.execute(
            """UPDATE jobs SET state='queued',lease_owner=NULL,lease_expires_at=NULL,updated_at=?
               WHERE state='running' AND lease_expires_at<?""", (stamp, stamp)
        )
        row = db.execute(
            """SELECT j.* FROM jobs j JOIN vaults v ON v.id=j.vault_id
               WHERE j.state='queued' AND j.not_before<=? AND v.paused=0
               ORDER BY COALESCE(v.last_job_claim_at,0),j.priority,j.created_at LIMIT 1""", (stamp,)
        ).fetchone()
        if not row:
            return None
        db.execute(
            """UPDATE jobs SET state='running',attempts=attempts+1,lease_owner=?,lease_expires_at=?,updated_at=?
               WHERE id=? AND state='queued'""", (WORKER_ID, stamp + 1800, stamp, row["id"])
        )
        if db.execute("SELECT changes()").fetchone()[0] != 1:
            return None
        db.execute("UPDATE vaults SET last_job_claim_at=? WHERE id=?", (stamp, row["vault_id"]))
        return dict(row)


def finish_job(job_id: str) -> None:
    with connect() as db:
        db.execute("UPDATE jobs SET state='complete',lease_owner=NULL,lease_expires_at=NULL,updated_at=? WHERE id=?", (now(), job_id))


def fail_job(job: dict, exc: Exception) -> None:
    retry = job["attempts"] + 1 < job["max_attempts"]
    state = "queued" if retry else "failed"
    delay = min(300, 5 * 2 ** job["attempts"])
    message = f"{type(exc).__name__}: {exc}"[:2000]
    with connect() as db:
        db.execute(
            """UPDATE jobs SET state=?,not_before=?,lease_owner=NULL,lease_expires_at=NULL,updated_at=?,error=? WHERE id=?""",
            (state, now() + delay, now(), message, job["id"]),
        )
        if not retry and job["kind"] == "inspect_temporary":
            inspection_id = json.loads(job["payload_json"])["inspection_id"]
            row = db.execute("SELECT staging_path FROM temporary_inspections WHERE vault_id=? AND id=?", (job["vault_id"], inspection_id)).fetchone()
            if row and row["staging_path"]:
                Path(row["staging_path"]).unlink(missing_ok=True)
            db.execute(
                "UPDATE temporary_inspections SET state='failed',error=?,staging_path=NULL,completed_at=? WHERE vault_id=? AND id=?",
                (message, now(), job["vault_id"], inspection_id),
            )


def _replace_chunks(vault_id: str, document_id: str, version_id: str, extracted) -> list[dict]:
    stamp = now()
    records = []
    for ordinal, chunk in enumerate(extracted.chunks):
        chunk_id = "chk_" + uuid.uuid5(uuid.NAMESPACE_URL, f"{version_id}:{ordinal}:{chunk.text[:120]}").hex
        records.append({
            "id": chunk_id, "vault_id": vault_id, "document_id": document_id,
            "version_id": version_id, "ordinal": ordinal, "kind": chunk.kind,
            "heading_path": chunk.heading_path, "locator_json": json.dumps(chunk.locator, ensure_ascii=False),
            "text": chunk.text, "contextual_text": None, "token_estimate": max(1, len(chunk.text) // 4),
        })
    with connect() as db, transaction(db, immediate=True):
        db.execute("DELETE FROM chunks_fts WHERE vault_id=? AND version_id=?", (vault_id, version_id))
        db.execute("DELETE FROM chunks WHERE vault_id=? AND version_id=?", (vault_id, version_id))
        db.executemany(
            """INSERT INTO chunks(id,vault_id,document_id,version_id,ordinal,kind,heading_path,locator_json,text,contextual_text,token_estimate,created_at)
               VALUES(:id,:vault_id,:document_id,:version_id,:ordinal,:kind,:heading_path,:locator_json,:text,:contextual_text,:token_estimate,:created_at)""",
            [{**record, "created_at": stamp} for record in records],
        )
        db.executemany(
            """INSERT INTO chunks_fts(text,contextual_text,heading_path,vault_id,chunk_id,document_id,version_id)
               VALUES(:text,'',:heading_path,:vault_id,:id,:document_id,:version_id)""", records,
        )
    return records


def _replace_table_cells(vault_id: str, document_id: str, version_id: str, raw_path: Path, display_name: str) -> None:
    with connect() as db:
        db.execute("DELETE FROM table_cells WHERE vault_id=? AND version_id=?", (vault_id, version_id))
        if Path(display_name).suffix.casefold() != ".xlsx":
            return
        workbook = load_workbook(raw_path, read_only=True, data_only=False)
        rows = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    formula = str(cell.value) if cell.data_type == "f" else None
                    numeric = float(cell.value) if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) else None
                    rows.append((vault_id, document_id, version_id, sheet.title, cell.coordinate, cell.row, cell.column, str(cell.value), numeric, formula))
        db.executemany(
            """INSERT INTO table_cells(vault_id,document_id,version_id,sheet_name,cell_ref,row_number,column_number,raw_value,numeric_value,formula)
               VALUES(?,?,?,?,?,?,?,?,?,?)""", rows,
        )


def process_extract(job: dict) -> None:
    payload = json.loads(job["payload_json"])
    with connect() as db:
        row = db.execute(
            """SELECT v.*,d.display_name,d.size_bytes FROM versions v JOIN documents d
               ON d.vault_id=v.vault_id AND d.id=v.document_id WHERE v.vault_id=? AND v.id=? AND d.id=?""",
            (job["vault_id"], payload["version_id"], payload["document_id"]),
        ).fetchone()
    if not row or not row["raw_staging_path"]:
        raise RuntimeError("Staged source is no longer available; reconnect the source to retry")
    raw_path = Path(row["raw_staging_path"])
    preview_dir = settings().preview_dir / job["vault_id"] / payload["version_id"]
    extracted = extract_document(raw_path, row["display_name"], preview_dir)
    records = _replace_chunks(job["vault_id"], payload["document_id"], payload["version_id"], extracted)
    _replace_table_cells(job["vault_id"], payload["document_id"], payload["version_id"], raw_path, row["display_name"])
    warnings = list(extracted.warnings)
    provider = provider_settings(job["vault_id"])
    try:
        for start in range(0, len(records), 64):
            batch = records[start:start + 64]
            vectors = embed_texts(
                job["vault_id"], [item["text"] for item in batch], "document",
                f"embed:{payload['version_id']}:{start}:{provider['embedding_model']}",
            )
            upsert_text_vectors(job["vault_id"], batch, vectors, provider["embedding_dimensions"])
    except BudgetUnavailable as exc:
        warnings.append(f"Dense indexing pending: {exc}")
    visual_items = []
    visual_inputs = []
    for index, preview in enumerate(extracted.previews):
        if not records:
            break
        with Image.open(preview) as image:
            visual_inputs.append([[records[min(index, len(records)-1)]["text"][:500], image.convert("RGB").copy()]][0])
        visual_items.append({
            "id": f"visual:{payload['version_id']}:{index}", "chunk_id": records[min(index, len(records)-1)]["id"],
            "document_id": payload["document_id"], "version_id": payload["version_id"], "preview_path": str(preview),
        })
    if visual_items:
        try:
            for start in range(0, len(visual_items), 8):
                vectors = embed_visual_inputs(job["vault_id"], visual_inputs[start:start+8], "document", f"visual:{payload['version_id']}:{start}:{provider['visual_model']}")
                upsert_visual_vectors(job["vault_id"], visual_items[start:start+8], vectors, provider["embedding_dimensions"])
        except BudgetUnavailable as exc:
            warnings.append(f"Visual indexing pending: {exc}")
    with connect() as db, transaction(db, immediate=True):
        db.execute("DELETE FROM previews WHERE vault_id=? AND version_id=?", (job["vault_id"], payload["version_id"]))
        db.executemany(
            """INSERT INTO previews(id,vault_id,document_id,version_id,ordinal,path,media_type,created_at)
               VALUES(?,?,?,?,?,?,?,?)""",
            [(f"preview_{uuid.uuid5(uuid.NAMESPACE_URL, str(path)).hex}", job["vault_id"], payload["document_id"], payload["version_id"], index, str(path), "image/" + ("webp" if path.suffix.casefold() == ".webp" else "jpeg"), now()) for index, path in enumerate(extracted.previews)],
        )
    raw_path.unlink(missing_ok=True)
    text_bytes = sum(len(item["text"].encode()) for item in records)
    preview_bytes = sum(path.stat().st_size for path in extracted.previews if path.exists())
    with connect() as db, transaction(db, immediate=True):
        db.execute(
            """UPDATE versions SET indexed_at=?,embedding_model=?,visual_model=?,raw_staging_path=NULL,
               text_bytes=?,preview_bytes=? WHERE vault_id=? AND id=?""",
            (now(), provider["embedding_model"], provider["visual_model"], text_bytes, preview_bytes, job["vault_id"], payload["version_id"]),
        )
        db.execute(
            """UPDATE documents SET current_version_id=?,state='indexed',error=? WHERE vault_id=? AND id=?""",
            (payload["version_id"], "\n".join(warnings) or None, job["vault_id"], payload["document_id"]),
        )
        obsolete = db.execute(
            """SELECT id,raw_staging_path FROM versions WHERE vault_id=? AND document_id=?
               AND id<>? AND indexed_at IS NULL""",
            (job["vault_id"], payload["document_id"], payload["version_id"]),
        ).fetchall()
        for old in obsolete:
            if old["raw_staging_path"]:
                Path(old["raw_staging_path"]).unlink(missing_ok=True)
            db.execute("DELETE FROM versions WHERE vault_id=? AND id=?", (job["vault_id"], old["id"]))
        db.execute(
            """DELETE FROM jobs WHERE vault_id=? AND kind='extract' AND state='failed'
               AND json_extract(payload_json,'$.document_id')=?""",
            (job["vault_id"], payload["document_id"]),
        )


def process_temporary_inspection(job: dict) -> None:
    inspection_id = json.loads(job["payload_json"])["inspection_id"]
    with connect() as db:
        row = db.execute(
            "SELECT * FROM temporary_inspections WHERE vault_id=? AND id=? AND state='queued'",
            (job["vault_id"], inspection_id),
        ).fetchone()
    if not row or not row["staging_path"]:
        raise RuntimeError("Temporary inspection expired or is unavailable")
    raw_path = Path(row["staging_path"])
    preview_dir = settings().preview_dir / "temporary" / inspection_id
    try:
        extracted = extract_document(raw_path, row["file_name"], preview_dir)
        full_text = "\n\n".join(chunk.text for chunk in extracted.chunks)
        truncated = len(full_text) > 16000
        warnings = [*extracted.warnings]
        if truncated:
            warnings.append("Temporary extracted context was truncated to 16,000 characters")
        with connect() as db:
            db.execute(
                """UPDATE temporary_inspections SET state='complete',extracted_text=?,warnings_json=?,staging_path=NULL,completed_at=?
                   WHERE vault_id=? AND id=?""",
                (full_text[:16000], json.dumps(warnings), now(), job["vault_id"], inspection_id),
            )
    finally:
        raw_path.unlink(missing_ok=True)
        for item in preview_dir.glob("*") if preview_dir.exists() else []:
            item.unlink(missing_ok=True)
        preview_dir.rmdir() if preview_dir.exists() else None


def process_job(job: dict) -> None:
    if job["kind"] == "extract":
        process_extract(job)
    elif job["kind"] == "inspect_temporary":
        process_temporary_inspection(job)
    elif job["kind"] == "delete_vectors":
        payload = json.loads(job["payload_json"])
        delete_document_vectors(job["vault_id"], payload["document_id"])
    elif job["kind"] == "cleanup":
        cleanup_expired()
    else:
        raise ValueError(f"Unknown job kind {job['kind']}")


def main() -> None:
    init_db()
    last_cleanup = 0.0
    while True:
        job = claim_job()
        if job:
            try:
                process_job(job)
                finish_job(job["id"])
            except Exception as exc:
                traceback.print_exc()
                fail_job(job, exc)
            continue
        if now() - last_cleanup > 3600:
            cleanup_expired()
            last_cleanup = now()
        time.sleep(settings().worker_poll_seconds)


if __name__ == "__main__":
    main()
