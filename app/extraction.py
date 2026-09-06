from __future__ import annotations

import json
import math
import os
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from .config import settings


@dataclass
class ExtractedChunk:
    text: str
    kind: str = "text"
    heading_path: str = ""
    locator: dict[str, Any] = field(default_factory=dict)


@dataclass
class ExtractedPreview:
    path: Path
    locator: dict[str, Any] = field(default_factory=dict)


@dataclass
class ExtractionResult:
    chunks: list[ExtractedChunk]
    previews: list[ExtractedPreview] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    coverage: dict[str, Any] = field(default_factory=dict)


def _split(text: str, locator: dict[str, Any], heading: str = "", kind: str = "text", limit: int = 1800) -> list[ExtractedChunk]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    paragraphs = re.split(r"\n\s*\n", text)
    output: list[ExtractedChunk] = []
    buffer = ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(buffer) + len(para) + 2 <= limit:
            buffer = f"{buffer}\n\n{para}".strip()
            continue
        if buffer:
            output.append(ExtractedChunk(buffer, kind, heading, dict(locator)))
        while len(para) > limit:
            cut = para.rfind(" ", 0, limit)
            cut = cut if cut > limit // 2 else limit
            output.append(ExtractedChunk(para[:cut].strip(), kind, heading, dict(locator)))
            para = para[cut:].strip()
        buffer = para
    if buffer:
        output.append(ExtractedChunk(buffer, kind, heading, dict(locator)))
    return output


def _docling(path: Path) -> ExtractionResult:
    from docling.document_converter import DocumentConverter

    result = DocumentConverter().convert(path)
    document = result.document
    chunks: list[ExtractedChunk] = []
    headings: list[str] = []
    for item, _level in document.iterate_items():
        text = getattr(item, "text", None)
        if not text and hasattr(item, "export_to_markdown"):
            try:
                text = item.export_to_markdown()
            except Exception:
                text = None
        if not text:
            continue
        label = str(getattr(item, "label", "text"))
        if "section_header" in label or label.endswith("title"):
            headings = [str(text).strip()]
        locator: dict[str, Any] = {}
        prov = getattr(item, "prov", None) or []
        pages = sorted({int(p.page_no) for p in prov if getattr(p, "page_no", None) is not None})
        if pages:
            locator["page"] = pages[0]
            if len(pages) > 1:
                locator["page_end"] = pages[-1]
        chunks.extend(_split(str(text), locator, " > ".join(headings), label))
    if not chunks:
        chunks = _split(document.export_to_markdown(), {}, kind="document")
    return ExtractionResult(chunks, coverage={"extractor": "docling", "complete_input": True})


def _text(path: Path) -> ExtractionResult:
    text = path.read_text(encoding="utf-8", errors="replace")
    chunks: list[ExtractedChunk] = []
    heading = ""
    for lineno, block in enumerate(re.split(r"(?m)(?=^#{1,6}\s)", text), 1):
        first = block.splitlines()[0] if block.splitlines() else ""
        if first.startswith("#"):
            heading = first.lstrip("# ").strip()
        chunks.extend(_split(block, {"line_group": lineno}, heading))
    return ExtractionResult(chunks, coverage={"extractor": "plain_text", "complete_input": True})


def _pdf_fallback(path: Path) -> ExtractionResult:
    chunks: list[ExtractedChunk] = []
    reader = PdfReader(path)
    for page_number, page in enumerate(reader.pages, 1):
        chunks.extend(_split(page.extract_text() or "", {"page": page_number}))
    return ExtractionResult(chunks, warnings=["Docling failed; used pypdf text fallback"], coverage={"pages": len(reader.pages)})


def _docx_fallback(path: Path) -> ExtractionResult:
    doc = WordDocument(path)
    chunks: list[ExtractedChunk] = []
    heading = ""
    buffer: list[str] = []
    for index, para in enumerate(doc.paragraphs):
        if para.style and para.style.name.startswith("Heading"):
            if buffer:
                chunks.extend(_split("\n".join(buffer), {"paragraph": index - len(buffer)}, heading))
                buffer = []
            heading = para.text.strip()
        buffer.append(para.text)
    chunks.extend(_split("\n".join(buffer), {"paragraph": max(0, len(doc.paragraphs) - len(buffer))}, heading))
    return ExtractionResult(chunks, warnings=["Docling failed; used DOCX fallback"], coverage={"paragraphs": len(doc.paragraphs)})


def _pptx_fallback(path: Path) -> ExtractionResult:
    prs = Presentation(path)
    chunks: list[ExtractedChunk] = []
    for slide_number, slide in enumerate(prs.slides, 1):
        text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
        chunks.extend(_split(text, {"slide": slide_number}, kind="slide"))
    return ExtractionResult(chunks, warnings=["Docling failed; used PPTX fallback"], coverage={"slides": len(prs.slides)})


def _xlsx_structured(path: Path, base: ExtractionResult | None = None) -> ExtractionResult:
    workbook = load_workbook(path, read_only=True, data_only=False)
    chunks = list(base.chunks if base else [])
    warnings = list(base.warnings if base else [])
    for sheet in workbook.worksheets:
        batch: list[str] = []
        start_row = 1
        for row_number, row in enumerate(sheet.iter_rows(values_only=False), 1):
            values = [cell.value for cell in row]
            rendered = " | ".join("" if value is None else str(value) for value in values)
            if rendered.strip(" |"):
                batch.append(rendered)
            if len(batch) >= 40:
                chunks.extend(_split("\n".join(batch), {"sheet": sheet.title, "row_start": start_row, "row_end": row_number}, sheet.title, "spreadsheet"))
                batch, start_row = [], row_number + 1
        if batch:
            chunks.extend(_split("\n".join(batch), {"sheet": sheet.title, "row_start": start_row, "row_end": sheet.max_row}, sheet.title, "spreadsheet"))
    warnings.append("Spreadsheet cells are retained with sheet/row locators; formulas are shown without recalculation")
    return ExtractionResult(chunks, base.previews if base else [], warnings, {"extractor": "docling+openpyxl" if base else "openpyxl", "sheets": workbook.sheetnames})


def _image_preview(path: Path, preview_dir: Path) -> list[ExtractedPreview]:
    preview_dir.mkdir(parents=True, exist_ok=True)
    target = preview_dir / "image.webp"
    with Image.open(path) as image:
        image.thumbnail((1800, 1800))
        image.convert("RGB").save(target, "WEBP", quality=82, method=4)
    return [ExtractedPreview(target, {"image": 1})]


def _image_ocr(path: Path) -> ExtractionResult:
    completed = subprocess.run(
        ["tesseract", str(path), "stdout", "-l", "eng+vie"],
        check=True,
        capture_output=True,
        text=True,
    )
    chunks = _split(completed.stdout, {"image": 1}, kind="ocr")
    return ExtractionResult(chunks, warnings=["Docling failed; used local Tesseract English/Vietnamese OCR"], coverage={"images": 1})


def _pdf_previews(path: Path, preview_dir: Path) -> tuple[list[ExtractedPreview], list[str]]:
    preview_dir.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(path)
    page_count = len(reader.pages)
    sampled = min(page_count, 120)
    subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "120", "-f", "1", "-l", str(sampled), str(path), str(preview_dir / "page")],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    warnings = [] if sampled == page_count else [f"Visual previews cover the first {sampled} of {page_count} PDF pages"]
    def page_number(item: Path) -> int:
        match = re.search(r"-(\d+)$", item.stem)
        if not match:
            raise RuntimeError(f"Rendered PDF preview has no page number: {item.name}")
        return int(match.group(1))

    previews = [ExtractedPreview(item, {"page": page_number(item)}) for item in sorted(preview_dir.glob("page-*.jpg"), key=page_number)]
    return previews, warnings


def _pdf_ocr(previews: list[ExtractedPreview]) -> list[ExtractedChunk]:
    chunks: list[ExtractedChunk] = []
    for preview in previews:
        completed = subprocess.run(
            ["tesseract", str(preview.path), "stdout", "-l", "eng+vie"],
            check=True,
            capture_output=True,
            text=True,
        )
        chunks.extend(_split(completed.stdout, dict(preview.locator), kind="ocr"))
    return chunks


def _media(path: Path, preview_dir: Path) -> ExtractionResult:
    from faster_whisper import WhisperModel

    preview_dir.mkdir(parents=True, exist_ok=True)
    is_video = path.suffix.casefold() in {".mp4", ".mov", ".mkv", ".webm"}
    audio_path = preview_dir / "audio.wav"
    warnings: list[str] = []
    chunks: list[ExtractedChunk] = []
    language = None
    duration = None
    try:
        subprocess.run(
            ["ffmpeg", "-nostdin", "-loglevel", "error", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", "-y", str(audio_path)],
            check=True,
        )
    except subprocess.CalledProcessError:
        if not is_video:
            raise
        warnings.append("Video has no usable audio stream; indexed visual frames only")
    if audio_path.is_file():
        model = WhisperModel("small", device="cpu", compute_type="int8", cpu_threads=2)
        segments, info = model.transcribe(str(audio_path), beam_size=3, vad_filter=True)
        language, duration = info.language, info.duration
        for segment in segments:
            chunks.extend(_split(segment.text, {"time_start": round(segment.start, 2), "time_end": round(segment.end, 2)}, kind="transcript"))
        audio_path.unlink(missing_ok=True)
    previews: list[ExtractedPreview] = []
    if is_video:
        pattern = preview_dir / "frame-%04d.jpg"
        subprocess.run(
            ["ffmpeg", "-nostdin", "-loglevel", "error", "-i", str(path), "-vf", f"fps=1/{settings().frame_interval_seconds},scale='min(1280,iw)':-2", "-frames:v", str(settings().max_video_frames), "-q:v", "4", "-y", str(pattern)],
            check=True,
        )
        previews = [
            ExtractedPreview(
                item,
                {"time_start": (index - 1) * settings().frame_interval_seconds, "time_end": index * settings().frame_interval_seconds},
            )
            for index, item in enumerate(sorted(preview_dir.glob("frame-*.jpg")), 1)
        ]
        if previews:
            warnings.append(f"Video frames sampled every {settings().frame_interval_seconds}s; fleeting events can be missed")
    return ExtractionResult(
        chunks,
        previews,
        warnings,
        {"language": language, "duration_seconds": duration, "frame_count": len(previews)},
    )


def extract_document(path: Path, display_name: str, preview_dir: Path) -> ExtractionResult:
    suffix = Path(display_name).suffix.casefold()
    if suffix in {".txt", ".md", ".markdown"}:
        return _text(path)
    if suffix in {".mp3", ".m4a", ".wav", ".flac", ".mp4", ".mov", ".mkv", ".webm"}:
        return _media(path, preview_dir)
    docling_result: ExtractionResult | None = None
    try:
        docling_result = _docling(path)
    except Exception as exc:
        docling_error = f"Docling extraction failed: {type(exc).__name__}: {exc}"
    else:
        docling_error = ""
    if suffix == ".xlsx":
        result = _xlsx_structured(path, docling_result)
    elif docling_result:
        result = docling_result
    elif suffix == ".pdf":
        result = _pdf_fallback(path)
    elif suffix == ".docx":
        result = _docx_fallback(path)
    elif suffix == ".pptx":
        result = _pptx_fallback(path)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
        result = _image_ocr(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    if docling_error:
        result.warnings.insert(0, docling_error)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
        result.previews.extend(_image_preview(path, preview_dir))
    elif suffix == ".pdf":
        previews, preview_warnings = _pdf_previews(path, preview_dir)
        result.previews.extend(previews)
        result.warnings.extend(preview_warnings)
        if not result.chunks:
            result.chunks = _pdf_ocr(previews)
            result.warnings.append("Structured extraction found no text; used local English/Vietnamese OCR on rendered pages")
    if not result.chunks and not result.previews:
        raise RuntimeError("Extraction completed without readable content")
    return result
