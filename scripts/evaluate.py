#!/usr/bin/env python3
"""Generate and exercise a small synthetic multilingual/multimodal corpus."""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import tempfile
import time
from pathlib import Path

import httpx
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


def corpus(root: Path) -> list[Path]:
    (root / "exact.md").write_text("# Operations\nThe immutable technical identifier is PVR-TECH-7Q9X. Never shorten it.\n", encoding="utf-8")
    (root / "vietnamese.txt").write_text("Ngân sách dự án Sen Vàng là 420 triệu đồng. Người phê duyệt là Lan.\n", encoding="utf-8")
    (root / "revision-old.md").write_text("# Historical draft\nThe Aurora launch was planned for 3 September 2026. This draft is superseded.\n", encoding="utf-8")
    (root / "revision-current.md").write_text("# Approved schedule\nThe approved Aurora launch date is 14 October 2026.\n", encoding="utf-8")
    (root / "multi-a.md").write_text("# North site\nThe North site owns sensor calibration under programme Juniper.\n", encoding="utf-8")
    (root / "multi-b.md").write_text("# South site\nThe South site owns network commissioning under programme Juniper.\n", encoding="utf-8")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Costs"
    sheet.append(["Item", "AUD"])
    for item, amount in (("Alpha", 10), ("Beta", 20), ("Gamma", 30)):
        sheet.append([item, amount])
    sheet["B5"] = "=SUM(B2:B4)"
    workbook.save(root / "costs.xlsx")

    doc = Document()
    doc.add_heading("Orchid decision", 1)
    doc.add_paragraph("The fallback colour for Orchid is cobalt blue.")
    doc.save(root / "decision.docx")

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Beacon checkpoint"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1))
    box.text_frame.text = "Checkpoint owner: Minh. Gate code: B-204."
    deck.save(root / "checkpoint.pptx")

    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((80, 100, 1120, 600), outline="navy", width=8)
    draw.text((140, 260), "SCANNED LABEL: KESTREL-88\nVisual marker: three orange circles", fill="black", spacing=20)
    draw.ellipse((700, 220, 780, 300), fill="orange")
    draw.ellipse((800, 220, 880, 300), fill="orange")
    draw.ellipse((900, 220, 980, 300), fill="orange")
    image.save(root / "scan.png")
    image.save(root / "scan.pdf", "PDF", resolution=120)

    subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-f", "lavfi", "-i", "flite=text=Audio checkpoint zephyr seven two:voice=slt", "-t", "4", "-y", str(root / "audio.wav")], check=True)
    subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=blue:s=640x360:d=5", "-f", "lavfi", "-i", "flite=text=Video marker indigo nine:voice=slt", "-shortest", "-c:v", "libx264", "-c:a", "aac", "-y", str(root / "video.mp4")], check=True)
    return sorted(path for path in root.iterdir() if path.is_file())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:5001")
    parser.add_argument("--token-file", required=True)
    parser.add_argument("--timeout", type=int, default=1800)
    args = parser.parse_args()
    token = Path(args.token_file).read_text().strip()
    headers = {"Authorization": f"Bearer {token}"}
    source_id = "src_synthetic_evaluation_2026"
    client = httpx.Client(base_url=args.base_url, headers=headers, timeout=300)
    client.post("/api/sources", json={"source_id": source_id, "name": "Synthetic evaluation", "platform": "linux", "root_label": "generated fixtures"}).raise_for_status()
    generation = 2026090601
    with tempfile.TemporaryDirectory(prefix="pvr-evaluation-") as temp:
        files = corpus(Path(temp))
        fixture_count = len(files)
        fixture_bytes = sum(path.stat().st_size for path in files)
        for path in files:
            raw = path.read_bytes()
            data = {
                "source_id": source_id,
                "relative_path": path.name,
                "display_name": path.name,
                "content_sha256": hashlib.sha256(raw).hexdigest(),
                "mtime_ns": "1700000000000000000",
                "size_bytes": str(len(raw)),
                "scan_generation": str(generation),
            }
            response = client.post("/api/ingest", data=data, files={"file": (path.name, raw)})
            response.raise_for_status()
        client.post("/api/reconcile", json={"source_id": source_id, "generation": generation, "complete": True}).raise_for_status()

    deadline = time.time() + args.timeout
    while time.time() < deadline:
        status = client.get("/api/status").json()
        active = sum(item["count"] for item in status["jobs"] if item["state"] in {"queued", "running"})
        if active == 0:
            break
        time.sleep(5)
    else:
        raise SystemExit("evaluation timed out waiting for queued work")

    checks = {}
    for name, question, expected_terms in (
        ("exact_identifier", "What is PVR-TECH-7Q9X?", ("PVR-TECH-7Q9X",)),
        ("vietnamese", "Ngân sách dự án Sen Vàng là bao nhiêu?", ("420 triệu đồng",)),
        ("conflict", "What are the conflicting Aurora launch dates?", ("3 September 2026", "14 October 2026")),
        ("multi_document", "Who owns work under programme Juniper?", ("North site", "South site")),
    ):
        result = client.post("/api/search", json={"question": question, "limit": 8}).json()
        joined = " ".join(item["preview"] for item in result["results"])
        complete_citations = bool(result["results"]) and all(
            item["document_id"] and item["version_id"] and item["chunk_id"] and item["locator"] is not None
            for item in result["results"]
        )
        checks[name] = {
            "passed": all(term.casefold() in joined.casefold() for term in expected_terms) and complete_citations,
            "required_terms": expected_terms,
            "result_count": len(result["results"]),
            "citations_have_ids_and_locators": complete_citations,
        }
    absent = client.post("/api/search", json={"question": "MARIGOLD-UNLISTED-404", "limit": 5}).json()
    absent_text = " ".join(item["preview"] for item in absent["results"])
    explicit_caution = any(
        "No supporting evidence" in warning or "require relevance validation" in warning
        for warning in absent["warnings"]
    )
    checks["unanswerable"] = {
        "passed": "MARIGOLD-UNLISTED-404" not in absent_text and explicit_caution,
        "result_count": len(absent["results"]),
        "warnings": absent["warnings"],
        "note": "Nearest-neighbour results are allowed, but must be explicitly cautioned and cannot contain invented support.",
    }

    search_costs = client.post("/api/search", json={"question": "Alpha Beta Gamma AUD Costs", "limit": 5}).json()
    spreadsheet = next((item for item in search_costs["results"] if item["title"] == "costs.xlsx"), None)
    if spreadsheet:
        calculated = client.post("/api/table/calculate", json={"document_id": spreadsheet["document_id"], "sheet_name": "Costs", "cell_range": "B2:B4", "operation": "sum"}).json()
        checks["structured_calculation"] = {"passed": calculated.get("value") == 60, "result": calculated}
    else:
        checks["structured_calculation"] = {"passed": False, "error": "spreadsheet was not retrieved"}

    status = client.get("/api/status").json()
    report = {"fixture_files": fixture_count, "fixture_bytes": fixture_bytes, "checks": checks, "status": status}
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0 if all(value["passed"] for value in checks.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
